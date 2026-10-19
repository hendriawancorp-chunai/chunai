"""Baca dokumen panjang per halaman (PDF/DOCX/PPTX/TXT/MD) agar ringkasan bisa bersitasi nomor halaman.
PDF lewat pdftotext (poppler); DOCX dikonversi ke PDF (LibreOffice) sehingga nomor halaman mendekati tampilan Word; PPTX per slide."""
from __future__ import annotations

import re
import shutil
import subprocess
from difflib import SequenceMatcher
from pathlib import Path

MAX_BYTES = 60 * 1024 * 1024
MAX_PAGES = 3000
TEXT_PAGE_CHARS = 3000
SUPPORTED = (".pdf", ".docx", ".pptx", ".txt", ".md")
_CACHE: dict[tuple, "Doc"] = {}


class DocReadError(ValueError):
    ...


class Doc:
    def __init__(self, name: str, pages: list[str], unit: str, note: str = ""):
        self.name, self.pages, self.unit, self.note = name, pages, unit, note

    @property
    def n(self) -> int:
        return len(self.pages)

    def empty_pages(self) -> list[int]:
        return [i for i, p in enumerate(self.pages, 1) if len(p.strip()) < 20]


def _pdf_pages(pdf: Path) -> list[str]:
    if not shutil.which("pdftotext"):
        raise DocReadError("pdftotext (poppler-utils) tidak tersedia. Pasang: sudo apt install poppler-utils")
    r = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, timeout=300)
    if r.returncode != 0:
        raise DocReadError("PDF tidak bisa dibaca (rusak atau terkunci kata sandi).")
    pages = r.stdout.decode("utf-8", "replace").split("\f")
    if pages and not pages[-1].strip():
        pages.pop()
    return pages


def _pptx_pages(path: Path) -> list[str]:
    from pptx import Presentation
    out = []
    for slide in Presentation(str(path)).slides:
        parts = []
        for sh in slide.shapes:
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                parts += [p.text for p in sh.text_frame.paragraphs if p.text.strip()]
            if getattr(sh, "has_table", False) and sh.has_table:
                parts += [" | ".join(c.text for c in row.cells) for row in sh.table.rows]
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                cats = list(ch.plots[0].categories) if len(ch.plots) else []
                parts.append(f"[grafik] kategori {cats}; " + "; ".join(f"{s.name}: {list(s.values)}" for s in ch.plots[0].series))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            n = slide.notes_slide.notes_text_frame.text.strip()
            if n:
                parts.append(f"[catatan pembicara] {n}")
        out.append("\n".join(parts))
    return out


def _split_text(text: str) -> list[str]:
    pages, cur = [], ""
    for para in re.split(r"\n{2,}", text):
        if cur and len(cur) + len(para) > TEXT_PAGE_CHARS:
            pages.append(cur)
            cur = ""
        cur += para + "\n\n"
    if cur.strip():
        pages.append(cur)
    return pages


def load(path: Path, cache_dir: Path | None = None) -> Doc:
    path = Path(path)
    if not path.is_file():
        raise DocReadError(f"File tidak ditemukan: {path.name}")
    if path.stat().st_size > MAX_BYTES:
        raise DocReadError(f"File terlalu besar (maks {MAX_BYTES // 1024 // 1024} MB).")
    key = (str(path), path.stat().st_mtime_ns, path.stat().st_size)
    if key in _CACHE:
        return _CACHE[key]
    ext = path.suffix.lower()
    note = ""
    if ext == ".pdf":
        pages, unit = _pdf_pages(path), "hlm"
    elif ext == ".pptx":
        pages, unit = _pptx_pages(path), "slide"
    elif ext == ".docx":
        from .office.render import to_pdf
        try:
            pdf = to_pdf(path, cache_dir or path.parent / ".cache")
            pages, unit = _pdf_pages(pdf), "hlm"
            note = "Nomor halaman DOCX dihitung dari konversi LibreOffice (bisa berbeda tipis dari tampilan Word)."
        except Exception:
            from docx import Document
            pages = _split_text("\n\n".join(p.text for p in Document(str(path)).paragraphs))
            unit, note = "bagian", "Konversi halaman gagal; nomor 'bagian' = potongan ±3000 karakter, BUKAN halaman Word."
    elif ext in (".txt", ".md"):
        pages, unit = _split_text(path.read_text(encoding="utf-8", errors="replace")), "bagian"
        note = "Berkas teks tidak punya halaman; 'bagian' = potongan ±3000 karakter."
    else:
        raise DocReadError(f"Format {ext or '(tanpa ekstensi)'} tidak didukung. Didukung: {', '.join(SUPPORTED)}.")
    if len(pages) > MAX_PAGES:
        raise DocReadError(f"Terlalu panjang ({len(pages)} halaman; maks {MAX_PAGES}). Pisahkan dulu.")
    doc = Doc(path.name, pages, unit, note)
    if len(_CACHE) > 20:
        _CACHE.clear()
    _CACHE[key] = doc
    return doc


def info(doc: Doc) -> str:
    chars = sum(len(p) for p in doc.pages)
    empty = doc.empty_pages()
    lines = [f"{doc.name}: {doc.n} {doc.unit}, ±{chars:,} karakter (≈{chars // 4:,} token)".replace(",", ".")]
    if doc.note:
        lines.append("Catatan: " + doc.note)
    if doc.n and len(empty) / doc.n > 0.5:
        lines.append(f"⚠ {len(empty)} dari {doc.n} {doc.unit} tidak berteks: kemungkinan HASIL PINDAI/gambar. OCR belum didukung; "
                     "isinya TIDAK bisa diringkas. Katakan terus terang kepada pengguna.")
    elif empty:
        lines.append(f"{doc.unit} tanpa teks (gambar/kosong): {empty[:30]}{'…' if len(empty) > 30 else ''}")
    per = max(1, round(chars / max(doc.n, 1)))
    lines.append(f"Rata-rata ±{per} karakter per {doc.unit}. Baca bertahap dengan baca_berkas (batas ±12.000 karakter per panggilan).")
    return "\n".join(lines)


def read_range(doc: Doc, start: int, end: int, max_chars: int = 12000) -> str:
    start, end = max(1, int(start)), min(doc.n, int(end))
    if start > doc.n or start > end:
        raise DocReadError(f"Rentang tidak valid; dokumen punya {doc.n} {doc.unit}.")
    out, used = [], 0
    for i in range(start, end + 1):
        txt = re.sub(r"[ \t]+\n", "\n", doc.pages[i - 1]).strip()
        block = f"=== {doc.unit} {i} ===\n{txt or '(tanpa teks)'}"
        if used and used + len(block) > max_chars:
            out.append(f"[BERHENTI sebelum {doc.unit} {i} karena batas panjang; lanjutkan dari {i}]")
            break
        out.append(block[:max_chars])
        used += len(block)
    return "\n\n".join(out)


def search(doc: Doc, query: str, max_hits: int = 25, width: int = 90) -> list[tuple[int, str]]:
    q = query.strip().lower()
    if not q:
        raise DocReadError("Kata pencarian kosong.")
    terms = [t for t in re.split(r"\s*\|\s*", q) if t]       # 'a | b' = salah satu
    hits = []
    for i, page in enumerate(doc.pages, 1):
        flat = " ".join(page.split())
        low = flat.lower()
        for t in terms:
            j = low.find(t)
            if j >= 0:
                hits.append((i, "…" + flat[max(0, j - width): j + len(t) + width] + "…"))
                break
        if len(hits) >= max_hits:
            break
    return hits


def _norm(s: str) -> str:
    s = s.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"').replace("­", "")
    s = re.sub(r"-\s*\n\s*", "", s)                          # kata terpotong di akhir baris
    return re.sub(r"\s+", " ", s).strip().lower()


def verify_quote(doc: Doc, page: int, quote: str) -> tuple[str, str]:
    """('ok'|'halaman_lain'|'mirip'|'tidak_ada', keterangan)."""
    q = _norm(quote).strip("\"'“” .…")
    if len(q) < 3:
        return "tidak_ada", "kutipan terlalu pendek untuk diverifikasi"
    if not 1 <= int(page) <= doc.n:
        return "tidak_ada", f"{doc.unit} {page} di luar dokumen (1-{doc.n})"
    if q in _norm(doc.pages[int(page) - 1]):
        return "ok", f"cocok di {doc.unit} {page}"
    others = [i for i, p in enumerate(doc.pages, 1) if i != int(page) and q in _norm(p)]
    if others:
        return "halaman_lain", f"teks ada, tetapi di {doc.unit} {', '.join(map(str, others[:5]))}, bukan {page}"
    cand = _norm(doc.pages[int(page) - 1])
    m = SequenceMatcher(None, cand, q, autojunk=False).find_longest_match(0, len(cand), 0, len(q))
    if len(q) >= 15 and m.size / len(q) >= 0.8:
        return "mirip", f"hampir sama tetapi bukan kutipan persis (potongan cocok {m.size}/{len(q)} karakter); jangan diberi tanda kutip"
    return "tidak_ada", f"tidak ditemukan di {doc.unit} {page} maupun halaman lain"
