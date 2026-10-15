"""Pembuat PPTX 16:9 gaya presentasi korporat: judul + subjudul berupa kalimat pesan, kartu KPI,
panel biru, tabel bersih, chart native (bisa diedit), proses/roadmap, footer bergradasi."""
from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .theme import Theme

W, H = 13.333, 7.5
_INLINE = re.compile(r"(\*\*[^*]+\*\*)")


def C(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr.upper())


def _strip_style(shape):
    st = shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if st is not None:
        shape._element.remove(st)


def fit_size(text: str | list[str], w: float, h: float, base: float, minimum: float = 9, line=1.22, cw=0.50) -> float:
    """Perkiraan ukuran font agar teks muat di kotak (w x h inci)."""
    paras = [text] if isinstance(text, str) else text
    size = base
    while size > minimum:
        lines = sum(max(1, math.ceil(len(p) * size * cw / 72 / max(w, 0.1))) for p in paras)
        if lines * size * line / 72 <= h:
            return size
        size -= 0.5
    return minimum


class Deck:
    def __init__(self, spec: dict, theme: Theme, base_dir: Path):
        self.spec, self.t, self.base = spec, theme, Path(base_dir)
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.warnings: list[str] = []
        self.no = 0

    # ---------- primitif ----------
    def rect(self, s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75):
        sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.shadow.inherit = False
        _strip_style(sh)
        if fill:
            sh.fill.solid(); sh.fill.fore_color.rgb = C(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = C(line); sh.line.width = Pt(lw)
        else:
            sh.line.fill.background()
        return sh

    def text(self, s, x, y, w, h, content, size=12, bold=False, color=None, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, italic=False, font=None, autofit=True, minimum=9, line_spacing=None, space_after=0):
        """content: str atau list[str] (paragraf). **tebal** didukung."""
        paras = [content] if isinstance(content, str) else list(content)
        if autofit:
            size = fit_size([re.sub(r"\*\*", "", p) for p in paras], w - 0.2, h - 0.1, size, minimum)
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.08); tf.margin_top = tf.margin_bottom = Inches(0.03)
        for i, ptxt in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            if line_spacing: p.line_spacing = line_spacing
            p.space_after = Pt(space_after)
            for part in _INLINE.split(ptxt):
                if not part:
                    continue
                b = bold
                if part.startswith("**") and part.endswith("**"):
                    part, b = part[2:-2], True
                r = p.add_run(); r.text = part
                r.font.size = Pt(size); r.font.bold = b; r.font.italic = italic
                r.font.name = font or self.t.font_slide
                r.font.color.rgb = C(color or self.t.text)
        return tb

    def bullets(self, s, x, y, w, h, items, size=12, color=None, dot=None, gap=5):
        dot = dot or self.t.primary
        size = fit_size([i + "    " for i in items], w - 0.35, h - 0.05, size, 9, line=1.22 + gap / 40)
        yy = y
        step = h / max(len(items), 1)
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.08); tf.margin_top = Inches(0.02)
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(gap)
            r = p.add_run(); r.text = "• "; r.font.size = Pt(size); r.font.bold = True
            r.font.color.rgb = C(dot); r.font.name = self.t.font_slide
            for part in _INLINE.split(it):
                if not part:
                    continue
                b = False
                if part.startswith("**") and part.endswith("**"):
                    part, b = part[2:-2], True
                r = p.add_run(); r.text = part; r.font.size = Pt(size); r.font.bold = b
                r.font.name = self.t.font_slide; r.font.color.rgb = C(color or self.t.text)
        return tb

    def gradient_bar(self, s, x, y, w, h):
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.shadow.inherit = False; sh.line.fill.background(); _strip_style(sh)
        sh.fill.gradient(); sh.fill.gradient_angle = 0
        st = sh.fill.gradient_stops
        st[0].color.rgb = C(self.t.green); st[0].position = 0
        st[1].color.rgb = C(self.t.primary); st[1].position = 1
        return sh

    # ---------- rangka ----------
    def frame(self, s, title: str, subtitle: str | None):
        self.no += 1
        self.rect(s, 0.17, 0.2, 0.06, 0.62, self.t.green)
        self.rect(s, 0.30, 0.2, 0.05, 0.62, self.t.primary)
        self.text(s, 0.5, 0.14, 10.8, 0.55, title, size=26, bold=True, color="0B1F3A", autofit=True, minimum=18)
        if subtitle:
            self.text(s, 0.5, 0.68, 10.9, 0.62, subtitle, size=13.5, color=self.t.primary, minimum=10)
        for key, x, al in (("logo_left", 0.5, None), ("logo_right", 11.6, None)):
            p = getattr(self.t, key)
            if p and Path(p).exists():
                s.shapes.add_picture(p, Inches(x if key == "logo_left" else 11.7), Inches(0.12), height=Inches(0.7))
        self.gradient_bar(s, 0, 7.18, W, 0.32)
        if self.t.footer_left:
            self.text(s, 0.15, 7.15, 8, 0.3, self.t.footer_left, size=8, italic=True, color="FFFFFF", autofit=False)
        right = f"{self.t.footer_right}   |   {self.no}" if self.t.footer_right else str(self.no)
        self.text(s, 8.6, 7.15, 4.6, 0.3, right, size=10, bold=True, color="FFFFFF", align=PP_ALIGN.RIGHT, autofit=False)

    def new(self, title, subtitle=None):
        s = self.prs.slides.add_slide(self.blank)
        self.frame(s, title, subtitle)
        return s

    # ---------- jenis slide ----------
    def cover(self, d):
        s = self.prs.slides.add_slide(self.blank); self.no += 1
        self.rect(s, 6.9, 0, W - 6.9, H, self.t.navy)
        self.rect(s, 0, 0, 6.9, H, "FFFFFF")
        for key, x in (("logo_left", 0.5), ("logo_right", 2.6)):
            p = getattr(self.t, key)
            if p and Path(p).exists():
                s.shapes.add_picture(p, Inches(x), Inches(0.3), height=Inches(0.6))
        if d.get("kicker"):
            self.text(s, 0.6, 1.7, 6, 0.6, d["kicker"], size=26, italic=True, bold=True, color=self.t.yellow, autofit=False)
        self.text(s, 0.6, 2.25, 6.2, 2.4, d["title"], size=34, bold=True, color=self.t.navy, minimum=22)
        y = 4.85
        for i, pill in enumerate(d.get("pills", [])[:3]):
            fill = (self.t.cyan, self.t.green, self.t.primary)[i]
            box = self.rect(s, 0.6, y, 5.8, 0.5, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            box.adjustments[0] = 0.5
            self.text(s, 0.75, y + 0.02, 5.5, 0.46, pill, size=14, bold=True, color="FFFFFF", anchor=MSO_ANCHOR.MIDDLE, minimum=10)
            y += 0.62
        self.text(s, 0.6, 6.25, 6, 0.4, d.get("label", ""), size=16, bold=True, italic=True, color=self.t.navy, autofit=False)
        self.text(s, 0.6, 6.62, 6, 0.55, [d.get("org", ""), d.get("year", "")], size=11, color=self.t.muted, autofit=False)
        if d.get("tagline"):
            self.text(s, 7.4, 6.55, 5.6, 0.6, d["tagline"], size=13, bold=True, color="FFFFFF", align=PP_ALIGN.LEFT, autofit=False)
        # motif titik
        for i in range(6):
            for j in range(4):
                self.rect(s, 7.4 + i * 0.28, 0.5 + j * 0.28, 0.06, 0.06, "3B6EA8", shape=MSO_SHAPE.OVAL)

    def ringkasan(self, d):
        s = self.new(d["title"], d.get("subtitle"))
        kp = d.get("kpis", [])
        if kp:
            wk = (W - 0.9 - 0.15 * (len(kp) - 1)) / len(kp)
            for i, k in enumerate(kp):
                x = 0.5 + i * (wk + 0.15)
                self.rect(s, x, 1.55, wk, 0.95, self.t.light)
                self.text(s, x + 0.05, 1.58, wk - 0.1, 0.5, k["v"], size=24, bold=True, color=self.t.primary, minimum=14)
                self.text(s, x + 0.05, 2.08, wk - 0.1, 0.4, k["l"], size=10.5, color="2B3A55", minimum=8)
        y0 = 2.7 if kp else 1.6
        hh = 6.95 - y0
        cards = d.get("cards", [])
        cw = 6.55 if d.get("panel") else W - 1.0
        if cards:
            ch = min((hh - 0.15 * (len(cards) - 1)) / len(cards), 1.9 if not d.get("panel") else 99)
            for i, c in enumerate(cards):
                y = y0 + i * (ch + 0.15)
                self.rect(s, 0.5, y, cw, ch, "F2F6FB"); self.rect(s, 0.5, y, 0.06, ch, self.t.primary)
                self.text(s, 0.7, y + 0.08, cw - 0.35, 0.4, c["h"], size=18, bold=True, color=self.t.navy, minimum=12)
                self.text(s, 0.7, y + 0.5, cw - 0.35, ch - 0.55, c["t"], size=16, color=self.t.text, minimum=10)
        pn = d.get("panel")
        if pn:
            self.rect(s, 7.3, y0, 5.55, hh, self.t.primary)
            self.text(s, 7.55, y0 + 0.15, 5.1, 0.5, pn["h"], size=17, bold=True, color="FFFFFF", minimum=12)
            items = pn.get("items", [])
            ih = (hh - 1.45) / max(len(items), 1)
            for i, it in enumerate(items):
                y = y0 + 0.8 + i * ih
                o = self.rect(s, 7.55, y + 0.05, 0.62, 0.62, "FFFFFF", shape=MSO_SHAPE.OVAL)
                self.text(s, 7.55, y + 0.12, 0.62, 0.48, it["tag"], size=14, bold=True, color=self.t.primary, align=PP_ALIGN.CENTER, autofit=False)
                self.text(s, 8.3, y, 4.45, 0.4, it["h"], size=14.5, bold=True, color="FFFFFF", minimum=10)
                if it.get("s"):
                    self.text(s, 8.3, y + 0.36, 4.45, 0.3, it["s"], size=12, color="DCEBFA", minimum=8)
            if pn.get("note"):
                self.text(s, 7.55, y0 + hh - 0.62, 5.1, 0.55, pn["note"], size=10.5, color="DCEBFA", minimum=8)

    def kolom(self, d):
        s = self.new(d["title"], d.get("subtitle"))
        cols = d["columns"]
        n = len(cols); gap = 0.18
        wc = (W - 1.0 - gap * (n - 1)) / n
        top = 1.55; bottom = 6.35 if d.get("band") else 6.95
        palette = [self.t.navy, self.t.primary, self.t.cyan, self.t.green]
        for i, c in enumerate(cols):
            x = 0.5 + i * (wc + gap)
            head = c.get("color") or palette[i % 4]
            self.rect(s, x, top, wc, 0.95, head)
            self.text(s, x + 0.1, top + 0.05, wc - 0.2, 0.4, c["h"], size=15, bold=True, color="FFFFFF", minimum=11)
            if c.get("sub"):
                self.text(s, x + 0.1, top + 0.45, wc - 0.2, 0.48, c["sub"], size=12, color="FFFFFF", minimum=8)
            self.rect(s, x, top + 0.95, wc, bottom - top - 0.95, "F2F6FB")
            if c.get("items"):
                self.bullets(s, x + 0.1, top + 1.12, wc - 0.2, bottom - top - 1.25, c["items"], size=16, gap=8)
        if d.get("band"):
            self.rect(s, 0.5, 6.48, W - 1.0, 0.5, "FFFFFF", line="C9D3E0")
            self.text(s, 0.6, 6.5, W - 1.2, 0.46, d["band"], size=12, bold=True, color=self.t.navy, align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE, minimum=9)

    def tabel(self, d):
        s = self.new(d["title"], d.get("subtitle"))
        self._table(s, d, 0.5, 1.6, W - 1.0, 5.2 if d.get("note") else 5.45)
        if d.get("note"):
            self.text(s, 0.5, 6.78, W - 1.0, 0.4, d["note"], size=10.5, color=self.t.primary, minimum=8)

    def _table(self, s, d, x, y, w, h, size=None):
        header, rows = d["header"], d["rows"]
        if size is None:
            size = 14 if len(rows) <= 4 else 12.5 if len(rows) <= 7 else 11 if len(rows) <= 10 else 10
        header_h = 0.5
        widths = d.get("widths") or [1] * len(header)
        tot = sum(widths)
        rh = min(0.85, (h - header_h) / max(len(rows), 1))
        shape = s.shapes.add_table(len(rows) + 1, len(header), Inches(x), Inches(y), Inches(w), Inches(header_h + rh * len(rows)))
        tbl = shape.table
        tblPr = shape._element.graphic.graphicData.tbl.tblPr
        for attr in ("firstRow", "bandRow"):
            tblPr.set(attr, "0")
        for i, wd in enumerate(widths):
            tbl.columns[i].width = Inches(w * wd / tot)
        hl = set(d.get("highlight", []))
        hlcol = d.get("highlight_col")
        cell_w = [w * wd / tot for wd in widths]

        def put(cell, txt, sz, bold, color, fill):
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid() if fill else cell.fill.background()
            if fill: cell.fill.fore_color.rgb = C(fill)
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            for part in _INLINE.split(str(txt)):
                if not part: continue
                b = bold
                if part.startswith("**") and part.endswith("**"): part, b = part[2:-2], True
                r = p.add_run(); r.text = part; r.font.size = Pt(sz); r.font.bold = b
                r.font.name = self.t.font_slide; r.font.color.rgb = C(color)

        for j, hd in enumerate(header):
            put(tbl.cell(0, j), hd, size, True, self.t.navy, None)
        tbl.rows[0].height = Inches(header_h)
        for i, row in enumerate(rows, start=1):
            sz = size
            longest = max(len(str(v)) * size * 0.5 / 72 / max(cell_w[k] - 0.16, 0.1) for k, v in enumerate(row))
            if longest * size * 1.25 / 72 > rh - 0.05:
                sz = max(8, size - 1.5)
                if (i - 1) < 1:
                    self.warnings.append(f"Slide {self.no}: tabel padat; font dikecilkan ke {sz}pt.")
            tbl.rows[i].height = Inches(rh)
            fill = self.t.zebra if (i % 2 == 0 and not (i - 1) in hl) else None
            if (i - 1) in hl:
                fill = "DCEBFA"
            for j, v in enumerate(row):
                f = "DCEBFA" if hlcol is not None and j == hlcol else fill
                put(tbl.cell(i, j), v, sz, j == 0 or (hlcol is not None and j == hlcol), self.t.text, f)
        # garis bawah header
        ln = s.shapes.add_connector(1, Inches(x), Inches(y + header_h), Inches(x + w), Inches(y + header_h))
        ln.line.color.rgb = C(self.t.primary); ln.line.width = Pt(1.5); _strip_style(ln)

    def grafik(self, d):
        s = self.new(d["title"], d.get("subtitle"))
        kp = d.get("kpis", [])
        top = 1.5
        if kp:
            wk = (W - 1.0 - 0.15 * (len(kp) - 1)) / len(kp)
            for i, k in enumerate(kp):
                x = 0.5 + i * (wk + 0.15)
                self.rect(s, x, top, wk, 0.8, self.t.light)
                self.text(s, x + 0.05, top + 0.02, wk - 0.1, 0.42, k["v"], size=20, bold=True, color=self.t.primary, minimum=12)
                self.text(s, x + 0.05, top + 0.43, wk - 0.1, 0.34, k["l"], size=10, color="2B3A55", minimum=8)
            top += 0.95
        ch = d["chart"]
        cd = CategoryChartData(); cd.categories = ch["categories"]
        for se in ch["series"]:
            cd.add_series(se["name"], se["values"])
        kind = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
                "barh": XL_CHART_TYPE.BAR_CLUSTERED}[ch.get("type", "bar")]
        has_side = bool(d.get("side_table"))
        cw = 5.2 if has_side else W - 1.0
        gf = s.shapes.add_chart(kind, Inches(0.5), Inches(top + (0.35 if ch.get("title") else 0)), Inches(cw),
                                Inches(6.75 - top - (0.35 if ch.get("title") else 0)), cd)
        if ch.get("title"):
            self.text(s, 0.5, top, cw, 0.35, ch["title"], size=13, bold=True, color=self.t.navy, autofit=False)
        c = gf.chart
        c.has_title = False
        c.has_legend = len(ch["series"]) > 1
        c.font.size = Pt(11); c.font.name = self.t.font_slide
        plot = c.plots[0]
        plot.has_data_labels = True
        dl = plot.data_labels; dl.font.size = Pt(12); dl.font.bold = True
        dl.number_format = ch.get("number_format", "General"); dl.number_format_is_linked = False
        if kind != XL_CHART_TYPE.LINE_MARKERS:
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
            plot.gap_width = 60
        c.value_axis.visible = False
        c.value_axis.has_major_gridlines = False
        c.category_axis.format.line.color.rgb = C("C9D3E0")
        for si, ser in enumerate(plot.series):
            col = (self.t.primary, self.t.green, self.t.cyan, self.t.navy)[si % 4]
            if kind == XL_CHART_TYPE.LINE_MARKERS:
                ser.format.line.color.rgb = C(col); ser.format.line.width = Pt(2.5); ser.smooth = False
            else:
                ser.format.fill.solid(); ser.format.fill.fore_color.rgb = C("A6A6A6" if "highlight" in ch else col)
                if "highlight" in ch:
                    for pi in range(len(ch["categories"])):
                        pt = ser.points[pi]; pt.format.fill.solid()
                        pt.format.fill.fore_color.rgb = C(self.t.primary if pi == ch["highlight"] else "A6A6A6")
        if has_side:
            self._table(s, d["side_table"], 6.0, top, W - 6.5, 6.75 - top, size=12.5)
        if d.get("note"):
            self.text(s, 0.5, 6.78, W - 1.0, 0.4, d["note"], size=10.5, color=self.t.muted, minimum=8)

    def proses(self, d):
        s = self.new(d["title"], d.get("subtitle"))
        steps = d["steps"]; n = len(steps)
        side = d.get("side")
        total_w = (8.55 if side else W - 1.0)
        gap = 0.28
        wc = (total_w - gap * (n - 1)) / n
        top = 1.55
        palette = [self.t.navy, self.t.primary, self.t.primary, self.t.cyan]
        for i, st in enumerate(steps):
            x = 0.5 + i * (wc + gap)
            self.rect(s, x, top, wc, 0.95, palette[min(i, 3)])
            self.text(s, x + 0.1, top + 0.04, wc - 0.2, 0.42, st["h"], size=15, bold=True, color="FFFFFF", minimum=11)
            if st.get("sub"):
                self.text(s, x + 0.1, top + 0.46, wc - 0.2, 0.46, st["sub"], size=10.5, color="DCEBFA", minimum=8)
            self.rect(s, x, top + 0.95, wc, 4.2 if not d.get("band") else 3.75, "F2F6FB")
            if st.get("lead"):
                self.text(s, x + 0.1, top + 1.02, wc - 0.2, 0.55, st["lead"], size=13, bold=True, color=self.t.navy, minimum=10)
            ly = top + (1.6 if st.get("lead") else 1.1)
            self.bullets(s, x + 0.1, ly, wc - 0.2, (top + 0.95 + (4.2 if not d.get("band") else 3.75)) - ly - 0.1, st.get("items", []), size=14, gap=7)
            if i < n - 1:
                ar = self.rect(s, x + wc + 0.02, top + 0.3, gap - 0.04, 0.35, self.t.cyan, shape=MSO_SHAPE.CHEVRON)
        if d.get("band"):
            self.rect(s, 0.5, 6.35, total_w, 0.5, "FFFFFF", line="C9D3E0")
            self.text(s, 0.6, 6.37, total_w - 0.2, 0.46, d["band"], size=12, bold=True, color=self.t.navy, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if side:
            self.rect(s, 9.3, top, 3.55, 5.4, self.t.primary)
            self.text(s, 9.5, top + 0.1, 3.2, 0.3, side.get("kicker", "HIGHLIGHT"), size=10, bold=True, color="CFE4F7", autofit=False)
            self.text(s, 9.5, top + 0.4, 3.2, 0.9, side["h"], size=16, bold=True, color="FFFFFF", minimum=11)
            stats = side.get("stats", []); sh = 3.4 / max(len(stats), 1)
            for i, st in enumerate(stats):
                y = top + 1.4 + i * sh
                self.text(s, 9.5, y, 1.0, 0.6, st["v"], size=28, bold=True, color="C5E58E", align=PP_ALIGN.CENTER, minimum=16)
                self.text(s, 10.55, y + 0.05, 2.2, 0.7, st["l"], size=11.5, color="FFFFFF", minimum=8)
            if side.get("note"):
                self.text(s, 9.5, top + 4.85, 3.2, 0.5, side["note"], size=9.5, italic=True, color="DCEBFA", minimum=7)

    def teks(self, d):
        s = self.new(d["title"], d.get("subtitle"))
        self.bullets(s, 0.6, 1.7, W - 1.2, 5.1, d["bullets"], size=22, gap=14)

    def penutup(self, d):
        s = self.new(d["title"], d.get("subtitle"))
        pw = 7.9 if d.get("highlight") else W - 1.0
        self.rect(s, 0.5, 1.6, pw, 5.2, "F2F6FB"); self.rect(s, 0.5, 1.6, 0.06, 5.2, self.t.primary)
        self.bullets(s, 0.8, 1.9, pw - 0.5, 4.7, d["bullets"], size=24 if not d.get("highlight") else 22, gap=18)
        if d.get("highlight"):
            self.rect(s, 8.7, 1.6, 4.15, 5.2, self.t.primary)
            self.text(s, 8.95, 1.8, 3.7, 4.8, d["highlight"], size=22, bold=True, color="FFFFFF", anchor=MSO_ANCHOR.MIDDLE, minimum=11)

    # ---------- orkestrasi ----------
    KINDS = {"cover", "ringkasan", "kolom", "tabel", "grafik", "proses", "teks", "penutup"}

    def build(self):
        for i, d in enumerate(self.spec.get("slides", []), 1):
            k = d.get("kind")
            if k not in self.KINDS:
                raise ValueError(f"Slide {i}: kind '{k}' tidak dikenal (pilihan: {sorted(self.KINDS)})")
            getattr(self, k)(d)
            if d.get("notes"):
                self.prs.slides[-1].notes_slide.notes_text_frame.text = d["notes"]
        return self.prs


def build_pptx(spec: dict, out_path: Path, theme: Theme, base_dir: Path) -> tuple[Path, list[str]]:
    deck = Deck(spec, theme, base_dir)
    prs = deck.build()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path, deck.warnings
