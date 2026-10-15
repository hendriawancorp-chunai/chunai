"""Pembuat DOCX gaya artikel/kajian: sampul, executive summary, bab bernomor, tabel bertema,
keterangan Tabel/Gambar otomatis, daftar pustaka, footer 'Page X of Y'."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .theme import Theme, rgb

_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
_LEAD_NUM = re.compile(r"^\s*\d+(\.\d+)*\.?\s+")
_LEAD_CAP = re.compile(r"^\s*(Tabel|Gambar|Table|Figure)\s+\d+\s*[.:\-–]?\s*", re.I)
_REF_TITLE = re.compile(r"^\s*(daftar\s+pustaka|daftar\s+rujukan|referensi|references|bibliography)\s*$", re.I)


def _shade(cell, hexcolor: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _cell_margins(table, top=50, bottom=50, left=80, right=80):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblPr.append(m)


def _borders(table, color="14284B", size=4):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(size)); e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)


def _no_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _field(paragraph, instr: str, size=8):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        elif text == instr:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = f" {instr} "
        else:
            el = OxmlElement("w:t"); el.text = text
        run._r.append(el)
    run.font.size = Pt(size); run.bold = True


def add_runs(par, text: str, size=None, color=None, font=None, bold=False, italic=False):
    for part in _INLINE.split(text):
        if not part:
            continue
        b, i = bold, italic
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], True
        r = par.add_run(part)
        r.bold, r.italic = b, i
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = rgb(color)
        if font:
            r.font.name = font
            r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    return par


class DocBuilder:
    def __init__(self, spec: dict, theme: Theme, base_dir: Path):
        self.spec, self.t, self.base = spec, theme, Path(base_dir)
        self.doc = Document()
        self.n = [0, 0, 0]
        self.tab_no = self.fig_no = 0
        self._styles()
        self._page()

    # ---------- dasar ----------
    def _styles(self):
        st = self.doc.styles
        n = st["Normal"]
        n.font.name, n.font.size = self.t.font_body, Pt(11)
        n.element.rPr.rFonts.set(qn("w:eastAsia"), self.t.font_body)
        n.font.color.rgb = rgb(self.t.text)
        pf = n.paragraph_format
        pf.space_after, pf.line_spacing = Pt(6), 1.15
        for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
            h = st[name]
            h.font.name, h.font.size, h.font.bold = self.t.font_heading, Pt(size), True
            h.font.color.rgb = rgb("000000")
            h.element.rPr.rFonts.set(qn("w:eastAsia"), self.t.font_heading)
            h.element.rPr.rFonts.set(qn("w:ascii"), self.t.font_heading)
            h.element.rPr.rFonts.set(qn("w:hAnsi"), self.t.font_heading)
            h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(14 if name == "Heading 1" else 10), Pt(4)
            h.paragraph_format.keep_with_next = True

    def _page(self):
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.5)
        s.top_margin, s.bottom_margin = Cm(2.5), Cm(2.3)
        s.different_first_page_header_footer = True
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.paragraph_format.space_after = Pt(2)
        pPr = fp._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "6"), ("space", "2"), ("color", "000000")):
            bot.set(qn(f"w:{k}"), v)
        bdr.append(bot); pPr.append(bdr)
        r = fp.add_run("Page "); r.italic = True; r.font.size = Pt(8)
        _field(fp, "PAGE")
        r = fp.add_run(" of "); r.italic = True; r.font.size = Pt(8)
        _field(fp, "NUMPAGES")
        fp2 = s.footer.add_paragraph()
        r = fp2.add_run(self.spec.get("footer", "")); r.italic = True; r.font.size = Pt(8)

    # ---------- sampul ----------
    def cover(self):
        c = self.spec.get("cover")
        if not c:
            return
        d = self.doc
        for key in ("logo_left", "logo_right"):
            p = getattr(self.t, key)
            if p and Path(p).exists():
                par = d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.LEFT if key == "logo_left" else WD_ALIGN_PARAGRAPH.RIGHT
                par.add_run().add_picture(p, height=Cm(1.6))
        for _ in range(7):
            d.add_paragraph()
        p = d.add_paragraph(); add_runs(p, c.get("kicker", "").upper(), size=44, color=self.t.navy, font=self.t.font_heading, bold=True)
        p.paragraph_format.space_after = Pt(2)
        for line in c.get("title_lines", [c.get("title", "")]):
            p = d.add_paragraph(); add_runs(p, line, size=22, color=self.t.cyan, font=self.t.font_heading, bold=True)
            p.paragraph_format.space_after = Pt(0)
        for _ in range(9):
            d.add_paragraph()
        p = d.add_paragraph(); add_runs(p, c.get("label", ""), size=16, color=self.t.primary, italic=True)
        self._rule(p)
        for lab, val in c.get("meta", []):
            p = d.add_paragraph(); p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(2.8))
            add_runs(p, f"{lab}\t: {val}", size=10.5, color=self.t.primary)
        for _ in range(3):
            d.add_paragraph()
        p = d.add_paragraph(); add_runs(p, c.get("org", ""), size=9, color=self.t.muted)
        d.add_page_break()

    def _rule(self, par):
        pPr = par._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "18"), ("space", "4"), ("color", self.t.primary)):
            b.set(qn(f"w:{k}"), v)
        bdr.append(b); pPr.append(bdr)

    # ---------- judul ----------
    def title_block(self):
        d, s = self.doc, self.spec
        if not s.get("title"):
            return
        for i, line in enumerate(s["title"].split("\n")):
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, line, size=16, bold=True, font=self.t.font_heading, color="111111")
            p.paragraph_format.space_after = Pt(0 if i < len(s["title"].split("\n")) - 1 else 12)
        for key, size, bold, ital in (("author", 11, True, False), ("affiliation", 10, False, False), ("contact", 9, False, True)):
            if key == "author" and str(s.get(key, "")).strip().lower() in ("chunai", "claude", "ai", "asisten"):
                continue                                              # jangan mengarang penulis; hanya nama yang diberikan pengguna
            if s.get(key):
                p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(p, s[key], size=size, bold=bold, italic=ital, color=self.t.navy if bold else self.t.muted)
                p.paragraph_format.space_after = Pt(2)
        d.add_paragraph()

    # ---------- blok ----------
    def heading(self, level: int, text: str):
        text = _LEAD_NUM.sub("", text) if level else text          # penomoran otomatis; buang nomor tulisan agen
        if level == 0:   # tanpa nomor, tengah (Executive Summary)
            p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, text, size=15, bold=True, font="Cambria", color=self.t.navy)
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(6), Pt(6)
            p.paragraph_format.keep_with_next = True
            return
        self.n[level - 1] += 1
        for i in range(level, 3):
            self.n[i] = 0
        num = ".".join(str(x) for x in self.n[:level]) + ("." if level == 1 else ".")
        h = self.doc.add_heading(level=level)
        h.paragraph_format.tab_stops.add_tab_stop(Cm(1.3 if level == 1 else 1.6 if level == 2 else 2.0))
        add_runs(h, f"{num}\t{text}", font=self.t.font_heading, color="000000", bold=True, size={1: 14, 2: 12, 3: 11}[level])

    def paragraph(self, text: str, indent=True):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent:
            p.paragraph_format.first_line_indent = Cm(1.0)
        add_runs(p, text)

    def bullets(self, items, numbered=False):
        for it in items:
            p = self.doc.add_paragraph(style="List Number" if numbered else "List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(1.4), Cm(-0.5)
            add_runs(p, it)

    def caption(self, text: str):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, text, size=9, italic=True, color=self.t.muted)
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(4), Pt(10)

    def table(self, t: dict):
        header, rows = t["header"], t["rows"]
        ncol = len(header)
        widths = t.get("widths") or [16 / ncol] * ncol
        tbl = self.doc.add_table(rows=1, cols=ncol)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        _borders(tbl); _cell_margins(tbl)
        for i, h in enumerate(header):
            c = tbl.rows[0].cells[i]; c.width = Cm(widths[i]); _shade(c, self.t.navy)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(c.paragraphs[0], h, size=9, bold=True, color="FFFFFF")
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
        _repeat_header(tbl.rows[0])
        for ri, row in enumerate(rows):
            r = tbl.add_row(); _no_split(r)
            for i, val in enumerate(row):
                c = r.cells[i]; c.width = Cm(widths[i])
                if ri % 2 == 1:
                    _shade(c, self.t.zebra)
                par = c.paragraphs[0]; par.paragraph_format.space_after = Pt(0); par.paragraph_format.line_spacing = 1.0
                add_runs(par, str(val), size=8.5, bold=(i == 0 and t.get("bold_first", True)))
        if t.get("caption"):
            self.tab_no += 1
            self.caption(f"Tabel {self.tab_no}. {_LEAD_CAP.sub('', t['caption'])}")
        else:
            self.doc.add_paragraph()

    def figure(self, f: dict):
        path = Path(f["path"])
        if not path.is_absolute():
            path = self.base / path
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(f.get("width_cm", 15)))
        if f.get("caption"):
            self.fig_no += 1
            self.caption(f"Gambar {self.fig_no}. {_LEAD_CAP.sub('', f['caption'])}")

    def references(self, refs: list[str]):
        if not refs:
            return
        self.doc.add_page_break()
        p = self.doc.add_paragraph()                                # judul tanpa nomor (gaya referensi)
        add_runs(p, self.spec.get("references_title", "Daftar Pustaka"), size=14, bold=True, font=self.t.font_heading, color=self.t.navy)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        for r in sorted(refs, key=str.lower):
            p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(1.0), Cm(-1.0)
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, r, size=10)

    # ---------- orkestrasi ----------
    def build(self) -> Document:
        self.cover()
        self.title_block()
        has_refs = bool(self.spec.get("references"))
        for b in self.spec.get("blocks", []):
            if "h" in b:
                if has_refs and _REF_TITLE.match(_LEAD_NUM.sub("", b["text"])):
                    continue                                          # Daftar Pustaka dibuat otomatis dari references
                self.heading(int(b["h"]), b["text"])
            elif "p" in b:
                self.paragraph(b["p"], indent=b.get("indent", True))
            elif "bullets" in b:
                self.bullets(b["bullets"])
            elif "numbered" in b:
                self.bullets(b["numbered"], numbered=True)
            elif "table" in b:
                self.table(b["table"])
            elif "figure" in b:
                self.figure(b["figure"])
            elif "pagebreak" in b:
                self.doc.add_page_break()
            else:
                raise ValueError(f"Blok tidak dikenal: {list(b)[:1]}")
        self.references(self.spec.get("references", []))
        return self.doc


def build_docx(spec: dict, out_path: Path, theme: Theme, base_dir: Path) -> Path:
    doc = DocBuilder(spec, theme, base_dir).build()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
