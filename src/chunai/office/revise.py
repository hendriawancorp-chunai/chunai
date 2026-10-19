"""Revisi dokumen yang sudah ada (DOCX/PPTX): baca struktur bernomor, lalu terapkan perubahan ke SALINAN baru.
Berkas asli tidak pernah ditimpa. Pemformatan dipertahankan (penggantian dilakukan pada level run)."""
from __future__ import annotations

import re
from pathlib import Path

MAX_OUTLINE = 14000
_TRUNC = 140


class ReviseError(ValueError):
    ...


def _cut(s: str, n: int = _TRUNC) -> str:
    s = " ".join(str(s).split())
    return s if len(s) <= n else s[:n] + f"…(+{len(s) - n})"


# ---------- penggantian teks yang mempertahankan format ----------
def replace_in_runs(runs, old: str, new: str) -> int:
    """Ganti semua `old` pada deret run (teks bisa terpecah di beberapa run). Format run pertama pada kecocokan dipertahankan."""
    if not old:
        return 0
    count, start = 0, 0
    while True:
        texts = [r.text for r in runs]
        full = "".join(texts)
        i = full.find(old, start)
        if i < 0:
            return count
        j = i + len(old)
        pos, a, b = 0, None, None
        for k, t in enumerate(texts):
            if a is None and i < pos + len(t):
                a, a_off = k, i - pos
            if j <= pos + len(t) and a is not None:
                b, b_off = k, j - pos
                break
            pos += len(t)
        if a is None or b is None:                       # kecocokan di luar run (tidak seharusnya)
            return count
        if a == b:
            runs[a].text = texts[a][:a_off] + new + texts[a][b_off:]
        else:
            runs[a].text = texts[a][:a_off] + new
            for k in range(a + 1, b):
                runs[k].text = ""
            runs[b].text = texts[b][b_off:]
        count += 1
        start = i + len(new)


def set_paragraph_text(par, text: str) -> None:
    runs = par.runs
    if not runs:
        par.add_run(text)
        return
    for br in par._p.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}br"):    # pptx: pemisah baris
        par._p.remove(br)
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


# =====================================================================
# DOCX
# =====================================================================
def _docx_ids(doc):
    paras = {f"P{i}": p for i, p in enumerate(doc.paragraphs, 1)}
    cells = {}
    for k, t in enumerate(doc.tables, 1):
        for r, row in enumerate(t.rows, 1):
            for c, cell in enumerate(row.cells, 1):
                cells[f"T{k}R{r}C{c}"] = cell
    return paras, cells


def _all_docx_paragraphs(doc):
    out = list(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                out += cell.paragraphs
    for s in doc.sections:
        for part in (s.header, s.footer):
            out += part.paragraphs
    seen, uniq = set(), []
    for p in out:                                         # sel gabungan muncul berulang
        if id(p._p) not in seen:
            seen.add(id(p._p))
            uniq.append(p)
    return uniq


def outline_docx(path: Path, cari: str = "") -> str:
    from docx import Document
    doc = Document(str(path))
    paras, cells = _docx_ids(doc)
    lines = [f"DOCX: {path.name} — {len(paras)} paragraf, {len(doc.tables)} tabel (header/footer hanya ikut penggantian global)"]
    q = cari.lower().strip()
    for pid, p in paras.items():
        text = p.text.strip()
        if not text:
            continue
        if q and q not in text.lower():
            continue
        style = p.style.name if p.style is not None else ""
        lines.append(f"{pid} [{style}] {_cut(text)}")
    for k, t in enumerate(doc.tables, 1):
        head = " | ".join(_cut(c.text, 30) for c in t.rows[0].cells) if t.rows else ""
        if q and not any(q in c.text.lower() for row in t.rows for c in row.cells):
            continue
        lines.append(f"T{k} ({len(t.rows)} baris x {len(t.columns)} kolom) judul kolom: {head}")
        for r, row in enumerate(t.rows, 1):
            if r > 25:
                lines.append(f"  … baris {r}+ tidak ditampilkan (id tetap T{k}R<n>C<m>)")
                break
            lines.append(f"  T{k}R{r}: " + " | ".join(_cut(c.text, 40) for c in row.cells))
    text = "\n".join(lines)
    return text if len(text) <= MAX_OUTLINE else text[:MAX_OUTLINE] + f"\n…(dipotong; pakai parameter cari='kata' untuk menyaring)"


def revise_docx(src: Path, dst: Path, ops: list[dict]) -> list[str]:
    from copy import deepcopy

    from docx import Document
    doc = Document(str(src))
    paras, cells = _docx_ids(doc)
    plan: list[tuple] = []
    errors: list[str] = []

    def par(op):
        pid = str(op.get("id", "")).strip().upper()
        if pid not in paras:
            errors.append(f"id paragraf '{op.get('id')}' tidak ada (pakai id dari baca_dokumen, mis. P12).")
        return paras.get(pid)

    for n, op in enumerate(ops, 1):
        kind = str(op.get("aksi", "")).strip().lower()
        if kind == "ganti_teks":
            old, new = str(op.get("cari", "")), str(op.get("ganti", ""))
            if not old:
                errors.append(f"op {n}: 'cari' kosong.")
                continue
            if op.get("id"):
                pid = str(op["id"]).strip().upper()
                target = paras.get(pid)
                cell = cells.get(pid)
                if target is None and cell is None:
                    errors.append(f"op {n}: id '{op['id']}' tidak ada.")
                    continue
                pars = [target] if target is not None else list(cell.paragraphs)
            else:
                pars = _all_docx_paragraphs(doc)
            if not any(old in p.text for p in pars):
                errors.append(f"op {n}: teks '{_cut(old, 50)}' tidak ditemukan{' di ' + str(op['id']) if op.get('id') else ''}.")
                continue
            plan.append(("ganti_teks", old, new, pars))
        elif kind == "ganti_paragraf":
            p = par(op)
            if p is not None:
                plan.append(("ganti_paragraf", p, str(op.get("teks", ""))))
        elif kind == "sisip_setelah":
            p = par(op)
            if p is not None:
                plan.append(("sisip_setelah", p, str(op.get("teks", ""))))
        elif kind == "hapus":
            p = par(op)
            if p is not None:
                plan.append(("hapus", p))
        elif kind == "ganti_sel":
            cid = str(op.get("id", "")).strip().upper()
            if cid not in cells:
                errors.append(f"op {n}: sel '{op.get('id')}' tidak ada (format T<tabel>R<baris>C<kolom>).")
            else:
                plan.append(("ganti_sel", cells[cid], str(op.get("teks", ""))))
        else:
            errors.append(f"op {n}: aksi '{op.get('aksi')}' tidak dikenal (ganti_teks, ganti_paragraf, sisip_setelah, hapus, ganti_sel).")
    if errors:
        raise ReviseError("Tidak ada yang diubah. Perbaiki dulu:\n- " + "\n- ".join(errors))

    report: list[str] = []
    for step in plan:
        if step[0] == "ganti_teks":
            _, old, new, pars = step
            total = sum(replace_in_runs(p.runs, old, new) for p in pars)
            report.append(f"ganti '{_cut(old, 40)}' → '{_cut(new, 40)}': {total} tempat")
        elif step[0] == "ganti_paragraf":
            _, p, text = step
            before = p.text
            set_paragraph_text(p, text)
            report.append(f"paragraf: '{_cut(before, 70)}' → '{_cut(text, 70)}'")
        elif step[0] == "sisip_setelah":
            _, p, text = step
            clone = deepcopy(p._p)
            p._p.addnext(clone)
            from docx.text.paragraph import Paragraph
            newp = Paragraph(clone, p._parent)
            set_paragraph_text(newp, text)
            report.append(f"sisip setelah '{_cut(p.text, 40)}': '{_cut(text, 70)}'")
        elif step[0] == "hapus":
            _, p = step
            report.append(f"hapus paragraf: '{_cut(p.text, 70)}'")
            p._p.getparent().remove(p._p)
        elif step[0] == "ganti_sel":
            _, cell, text = step
            before = cell.text
            ps = cell.paragraphs
            set_paragraph_text(ps[0], text)
            for extra in ps[1:]:
                set_paragraph_text(extra, "")
            report.append(f"sel: '{_cut(before, 50)}' → '{_cut(text, 50)}'")
    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dst))
    return report


# =====================================================================
# PPTX
# =====================================================================
def _walk(shapes, prefix=""):
    for i, sh in enumerate(shapes, 1):
        sid = f"{prefix}{i}"
        yield sid, sh
        if sh.shape_type == 6:                              # grup
            yield from _walk(sh.shapes, sid + ".")


def _pptx_ids(prs):
    out = {}
    for n, slide in enumerate(prs.slides, 1):
        for sid, sh in _walk(slide.shapes):
            out[f"S{n}.sh{sid}"] = sh
    return out


def _para_iter(sh):
    """Semua paragraf (kerangka teks dan sel tabel) pada satu bentuk."""
    if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
        yield from sh.text_frame.paragraphs
    if getattr(sh, "has_table", False) and sh.has_table:
        for row in sh.table.rows:
            for c in row.cells:
                yield from c.text_frame.paragraphs


def outline_pptx(path: Path, cari: str = "") -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    lines = [f"PPTX: {path.name} — {len(prs.slides)} slide"]
    q = cari.lower().strip()
    for n, slide in enumerate(prs.slides, 1):
        block = []
        for sid, sh in _walk(slide.shapes):
            key = f"S{n}.sh{sid}"
            if getattr(sh, "has_table", False) and sh.has_table:
                t = sh.table
                block.append(f"{key} [Tabel {len(t.rows)}x{len(t.columns)}]")
                for r, row in enumerate(t.rows, 1):
                    block.append(f"  {key}.r{r}: " + " | ".join(_cut(c.text, 30) for c in row.cells))
            elif getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                cats = list(ch.plots[0].categories) if len(ch.plots) else []
                block.append(f"{key} [Grafik {str(ch.chart_type).split('.')[-1].split(' ')[0]}] kategori: {cats}")
                for s in ch.plots[0].series:
                    block.append(f"  seri \"{s.name}\": {list(s.values)}")
            elif getattr(sh, "has_text_frame", False) and sh.has_text_frame and sh.text_frame.text.strip():
                block.append(f"{key} [{sh.name}] " + " ⏎ ".join(_cut(p.text, 110) for p in sh.text_frame.paragraphs if p.text.strip()))
        if q and not any(q in b.lower() for b in block):
            continue
        lines.append(f"--- Slide {n} ---")
        lines += block
    text = "\n".join(lines)
    return text if len(text) <= MAX_OUTLINE else text[:MAX_OUTLINE] + "\n…(dipotong; pakai parameter cari='kata' untuk menyaring)"


_CELL = re.compile(r"^(S\d+\.sh[\d.]+)\.r(\d+)c(\d+)$", re.I)


def revise_pptx(src: Path, dst: Path, ops: list[dict]) -> list[str]:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    prs = Presentation(str(src))
    shapes = _pptx_ids(prs)
    shapes_ci = {k.upper(): v for k, v in shapes.items()}
    plan: list[tuple] = []
    errors: list[str] = []

    for n, op in enumerate(ops, 1):
        kind = str(op.get("aksi", "")).strip().lower()
        if kind == "ganti_teks":
            old, new = str(op.get("cari", "")), str(op.get("ganti", ""))
            if not old:
                errors.append(f"op {n}: 'cari' kosong.")
                continue
            if op.get("slide"):
                try:
                    s = int(op["slide"])
                    prefix = f"S{s}."
                except (TypeError, ValueError):
                    errors.append(f"op {n}: 'slide' harus angka.")
                    continue
                pool = [sh for k, sh in shapes.items() if k.startswith(prefix)]
                if not 1 <= s <= len(prs.slides):
                    errors.append(f"op {n}: slide {s} tidak ada (1-{len(prs.slides)}).")
                    continue
            else:
                pool = list(shapes.values())
            pars = [p for sh in pool for p in _para_iter(sh)]
            if not any(old in p.text for p in pars):
                errors.append(f"op {n}: teks '{_cut(old, 50)}' tidak ditemukan{' di slide ' + str(op['slide']) if op.get('slide') else ''}.")
                continue
            plan.append(("ganti_teks", old, new, pars))
        elif kind == "ganti_shape_teks":
            sh = shapes_ci.get(str(op.get("id", "")).strip().upper())
            if sh is None or not getattr(sh, "has_text_frame", False) or not sh.has_text_frame:
                errors.append(f"op {n}: bentuk teks '{op.get('id')}' tidak ada (mis. S5.sh3).")
            else:
                plan.append(("ganti_shape_teks", sh, str(op.get("teks", ""))))
        elif kind == "ganti_sel":
            m = _CELL.match(str(op.get("id", "")).strip())
            sh = shapes_ci.get(m.group(1).upper()) if m else None
            if sh is None or not getattr(sh, "has_table", False) or not sh.has_table:
                errors.append(f"op {n}: sel '{op.get('id')}' tidak valid (format S5.sh7.r2c3).")
                continue
            r, c = int(m.group(2)), int(m.group(3))
            if not (1 <= r <= len(sh.table.rows) and 1 <= c <= len(sh.table.columns)):
                errors.append(f"op {n}: sel di luar tabel ({len(sh.table.rows)}x{len(sh.table.columns)}).")
                continue
            plan.append(("ganti_sel", sh.table.cell(r - 1, c - 1), str(op.get("teks", ""))))
        elif kind == "ganti_data_grafik":
            sh = shapes_ci.get(str(op.get("id", "")).strip().upper())
            series = op.get("seri")
            cats = op.get("kategori")
            if sh is None or not getattr(sh, "has_chart", False) or not sh.has_chart:
                errors.append(f"op {n}: grafik '{op.get('id')}' tidak ada.")
                continue
            old_cats = list(sh.chart.plots[0].categories)
            cats = list(cats) if cats else old_cats
            if not isinstance(series, list) or not series:
                errors.append(f"op {n}: 'seri' harus daftar [{{nama, nilai[]}}].")
                continue
            bad = [s for s in series if not isinstance(s, dict) or len(s.get("nilai", [])) != len(cats)]
            if bad:
                errors.append(f"op {n}: jumlah 'nilai' tiap seri harus {len(cats)} (sama dengan kategori).")
                continue
            plan.append(("ganti_data_grafik", sh, cats, series))
        else:
            errors.append(f"op {n}: aksi '{op.get('aksi')}' tidak dikenal (ganti_teks, ganti_shape_teks, ganti_sel, ganti_data_grafik).")
    if errors:
        raise ReviseError("Tidak ada yang diubah. Perbaiki dulu:\n- " + "\n- ".join(errors))

    report: list[str] = []
    for step in plan:
        if step[0] == "ganti_teks":
            _, old, new, pars = step
            total = sum(replace_in_runs(p.runs, old, new) for p in pars)
            report.append(f"ganti '{_cut(old, 40)}' → '{_cut(new, 40)}': {total} tempat")
        elif step[0] == "ganti_shape_teks":
            _, sh, text = step
            before = sh.text_frame.text
            lines = str(text).split("\n")
            ps = sh.text_frame.paragraphs
            for i, line in enumerate(lines):
                if i < len(ps):
                    set_paragraph_text(ps[i], line)
                else:                                        # tambah paragraf dengan format paragraf terakhir
                    from copy import deepcopy
                    clone = deepcopy(ps[-1]._p)
                    ps[-1]._p.addnext(clone)
                    ps = sh.text_frame.paragraphs
                    set_paragraph_text(ps[i], line)
            for extra in sh.text_frame.paragraphs[len(lines):]:
                extra._p.getparent().remove(extra._p)
            report.append(f"teks bentuk: '{_cut(before, 70)}' → '{_cut(text, 70)}'")
        elif step[0] == "ganti_sel":
            _, cell, text = step
            before = cell.text
            ps = cell.text_frame.paragraphs
            set_paragraph_text(ps[0], text)
            for extra in ps[1:]:
                extra._p.getparent().remove(extra._p)
            report.append(f"sel: '{_cut(before, 50)}' → '{_cut(text, 50)}'")
        elif step[0] == "ganti_data_grafik":
            _, sh, cats, series = step
            cd = CategoryChartData()
            cd.categories = cats
            for s in series:
                cd.add_series(str(s.get("nama", "")), [float(v) for v in s["nilai"]])
            sh.chart.replace_data(cd)
            report.append(f"data grafik diganti: {len(cats)} kategori, {len(series)} seri")
    dst.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dst))
    return report
