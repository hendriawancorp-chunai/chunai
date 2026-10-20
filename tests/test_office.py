import json
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation

from chunai.office.deckbuild import build_pptx, fit_size
from chunai.office.docbuild import build_docx
from chunai.office.doccheck import lint_doc_spec
from chunai.office.theme import Theme, load_theme

EX = Path(__file__).parent.parent / "examples"
DOC = json.loads((EX / "contoh_artikel.json").read_text())
DECK = json.loads((EX / "contoh_deck.json").read_text())


def test_example_doc_passes_lint():
    assert lint_doc_spec(DOC) == []


def test_lint_catches_problems():
    bad = json.loads(json.dumps(DOC))
    bad["blocks"].append({"p": "Pendapatan tumbuh 11 persen dan mencapai Rp4 triliun pada tahun lalu menurut data internal."})
    bad["blocks"].append({"p": "Klaim lain (Tidakada, 2020)."})
    bad["blocks"].append({"table": {"header": ["a"], "rows": [["x"]]}})
    issues = "\n".join(lint_doc_spec(bad))
    assert "tanpa sitasi" in issues and "Tidakada" in issues and "tanpa keterangan" in issues


def test_docx_structure(tmp_path):
    out = build_docx(DOC, tmp_path / "a.docx", Theme(), Path("."))
    d = Document(str(out))
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert heads[0].startswith("1.") and any(h.startswith("2.2.1.") for h in heads)
    assert any(p.text.startswith("Tabel 1.") for p in d.paragraphs)
    assert len(d.tables) == 1 and d.tables[0].rows[0].cells[0].text == "Dimensi"


def test_pptx_structure_and_native_chart(tmp_path):
    out, warns = build_pptx(DECK, tmp_path / "d.pptx", Theme(), Path("."))
    prs = Presentation(str(out))
    assert len(prs.slides) == len(DECK["slides"])
    assert any(sh.has_chart for sl in prs.slides for sh in sl.shapes)       # chart bisa diedit
    assert any(sh.has_table for sl in prs.slides for sh in sl.shapes)


def test_unknown_slide_kind_rejected(tmp_path):
    with pytest.raises(ValueError, match="tidak dikenal"):
        build_pptx({"slides": [{"kind": "ngawur"}]}, tmp_path / "x.pptx", Theme(), Path("."))


def test_fit_size_shrinks_long_text():
    assert fit_size("kata " * 200, 3, 1, 20) < 20
    assert fit_size("pendek", 5, 1, 20) == 20


def test_custom_theme_loads(tmp_path):
    (tmp_path / "x.json").write_text(json.dumps({"primary": "AA0000", "font_slide": "Poppins", "bukan_field": 1}))
    t = load_theme("x", tmp_path)
    assert t.primary == "AA0000" and t.font_slide == "Poppins"
    with pytest.raises(FileNotFoundError):
        load_theme("tidak-ada", tmp_path)


def test_ringkasan_and_penutup_adapt_without_optional_parts(tmp_path):
    spec = {"slides": [
        {"kind": "ringkasan", "title": "R", "cards": [{"h": "A", "t": "x"}, {"h": "B", "t": "y"}]},
        {"kind": "penutup", "title": "P", "bullets": ["satu", "dua"]}]}
    out, _ = build_pptx(spec, tmp_path / "s.pptx", Theme(), Path("."))
    prs = Presentation(str(out))
    widest = max(sh.width for sh in prs.slides[0].shapes if sh.shape_type == 1)       # kartu melebar penuh
    assert widest > 11_000_000
    assert len(prs.slides) == 2


def test_unknown_theme_error_lists_available(tmp_path):
    (tmp_path / "asdp.json").write_text("{}")
    with pytest.raises(FileNotFoundError, match="asdp"):
        load_theme("biru", tmp_path)


def test_docbuilder_cleans_agent_numbering_and_duplicate_references(tmp_path):
    spec = {"title": "T", "blocks": [
        {"h": 1, "text": "1. Pendahuluan"}, {"p": "Isi."},
        {"h": 2, "text": "2.1 Hasil Utama"},
        {"table": {"caption": "Tabel 1. Indikator", "header": ["a"], "rows": [["x"]]}},
        {"h": 1, "text": "Daftar Pustaka"}],
        "references": ["Contoh, A. (2026). Judul. https://contoh.example"]}
    out = build_docx(spec, tmp_path / "x.docx", Theme(), Path("."))
    d = Document(str(out))
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert heads == ["1.\tPendahuluan", "1.1.\tHasil Utama"]                 # tanpa nomor ganda; tanpa heading pustaka ganda
    assert any(p.text == "Tabel 1. Indikator" for p in d.paragraphs)         # tanpa "Tabel 1. Tabel 1."
    assert sum(p.text == "Daftar Pustaka" for p in d.paragraphs) == 1        # tepat satu, tanpa nomor


def test_kpi_only_summary_is_large(tmp_path):
    spec = {"slides": [{"kind": "ringkasan", "title": "R", "kpis": [{"v": "18 → 11", "l": "menit"}, {"v": "84%", "l": "puas"}, {"v": "6 → 9", "l": "loket"}]}]}
    out, _ = build_pptx(spec, tmp_path / "k.pptx", Theme(), Path("."))
    prs = Presentation(str(out))
    tiles = [sh for sh in prs.slides[0].shapes if sh.shape_type == 1 and sh.height > 914400 * 2]
    assert len(tiles) == 3                                                    # kotak KPI tinggi > 2 inci


def test_contact_sheet(tmp_path):
    from PIL import Image
    from chunai.office.tools import contact_sheet
    ps = []
    for i in range(3):
        p = tmp_path / f"{i}.png"; Image.new("RGB", (400, 225), "white").save(p); ps.append(p)
    out = contact_sheet(ps, tmp_path / "s.png")
    assert Image.open(out).size == (810, 460)


def test_fabricated_author_is_not_printed(tmp_path):
    spec = {"title": "T", "author": "ChunAI", "blocks": [{"p": "x"}]}
    d = Document(str(build_docx(spec, tmp_path / "a.docx", Theme(), Path("."))))
    assert all("ChunAI" not in p.text for p in d.paragraphs)
    spec["author"] = "Budi Santoso"
    d = Document(str(build_docx(spec, tmp_path / "b.docx", Theme(), Path("."))))
    assert any("Budi Santoso" in p.text for p in d.paragraphs)


def test_shrink_image_keeps_agent_reads_small(tmp_path):
    import numpy as np
    from PIL import Image

    from chunai.office.tools import shrink_image
    rng = np.random.default_rng(0)
    p = tmp_path / "besar.png"
    Image.fromarray(rng.integers(0, 255, (3000, 2400, 3), dtype=np.uint8)).save(p)          # derau acak: sulit dimampatkan (> 10 MB)
    assert p.stat().st_size > 5_000_000
    shrink_image(p)
    assert p.stat().st_size <= 700_000 and Image.open(p).width <= 1600


def test_agent_options_allow_large_messages(tmp_path):
    from chunai.agent import ChunAI
    a = ChunAI(None, workdir=tmp_path / "ws", audit_path=tmp_path / "a.jsonl")
    assert a._options().max_buffer_size >= 16 * 1024 * 1024
