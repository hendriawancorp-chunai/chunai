import json
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation

from chunai.office.deckbuild import build_pptx
from chunai.office.docbuild import build_docx
from chunai.office.revise import (ReviseError, outline_docx, outline_pptx, replace_in_runs, revise_docx, revise_pptx,
                                  set_paragraph_text)
from chunai.office.theme import Theme

EX = Path(__file__).parent.parent / "examples"


@pytest.fixture
def docx_file(tmp_path):
    return build_docx(json.loads((EX / "contoh_artikel.json").read_text()), tmp_path / "a.docx", Theme(), Path("."))


@pytest.fixture
def pptx_file(tmp_path):
    out, _ = build_pptx(json.loads((EX / "contoh_deck.json").read_text()), tmp_path / "d.pptx", Theme(), Path("."))
    return out


def test_replace_across_runs_keeps_format():
    d = Document()
    p = d.add_paragraph()
    p.add_run("Pendapatan ").bold = True
    p.add_run("tumbuh 11").italic = True
    p.add_run(" persen di 2026.")
    assert replace_in_runs(p.runs, "tumbuh 11 persen", "tumbuh 14 persen") == 1
    assert p.text == "Pendapatan tumbuh 14 persen di 2026."
    assert p.runs[0].bold and p.runs[1].italic                     # format run awal kecocokan tetap
    assert replace_in_runs(p.runs, "x", "y") == 0
    assert replace_in_runs(p.runs, "2026", "2026 dan 2026") == 1   # pengganti memuat teks lama: tidak berulang tanpa akhir


def test_outline_docx_has_ids(docx_file):
    o = outline_docx(docx_file)
    assert "\nP8 [" in o and "[Heading" in o and "T1 (" in o and "T1R1:" in o
    assert "Dimensi" in outline_docx(docx_file, cari="dimensi")


def test_docx_revision_writes_copy_and_reports(docx_file, tmp_path):
    before = Document(str(docx_file))
    pid = next(f"P{i}" for i, p in enumerate(before.paragraphs, 1) if p.style.name.startswith("Heading"))
    old_head = before.paragraphs[int(pid[1:]) - 1].text
    dst = tmp_path / "out" / "a_rev.docx"
    rep = revise_docx(docx_file, dst, [
        {"aksi": "ganti_paragraf", "id": pid, "teks": "JUDUL BARU"},
        {"aksi": "sisip_setelah", "id": pid, "teks": "Paragraf sisipan."},
        {"aksi": "ganti_sel", "id": "T1R1C1", "teks": "Aspek"},
    ])
    d = Document(str(dst))
    texts = [p.text for p in d.paragraphs]
    i = texts.index("JUDUL BARU")
    assert texts[i + 1] == "Paragraf sisipan." and d.paragraphs[i].style.name == d.paragraphs[i + 1].style.name
    assert d.tables[0].rows[0].cells[0].text == "Aspek"
    assert Document(str(docx_file)).paragraphs[int(pid[1:]) - 1].text == old_head       # asli utuh
    assert len(rep) == 3


def test_docx_global_replace_and_delete(docx_file, tmp_path):
    src_text = "\n".join(p.text for p in Document(str(docx_file)).paragraphs)
    word = next(w for w in src_text.split() if len(w) > 6 and src_text.count(w) >= 1)
    dst = tmp_path / "r.docx"
    rep = revise_docx(docx_file, dst, [{"aksi": "ganti_teks", "cari": word, "ganti": "XXXXXX"}])
    assert "XXXXXX" in "\n".join(p.text for p in Document(str(dst)).paragraphs) and "tempat" in rep[0]
    n = len(Document(str(docx_file)).paragraphs)
    revise_docx(docx_file, tmp_path / "h.docx", [{"aksi": "hapus", "id": f"P{n}"}])
    assert len(Document(str(tmp_path / "h.docx")).paragraphs) == n - 1


def test_docx_invalid_ops_change_nothing(docx_file, tmp_path):
    dst = tmp_path / "x.docx"
    with pytest.raises(ReviseError) as e:
        revise_docx(docx_file, dst, [{"aksi": "ganti_paragraf", "id": "P1", "teks": "ok"},
                                     {"aksi": "ganti_paragraf", "id": "P9999", "teks": "x"},
                                     {"aksi": "ganti_teks", "cari": "tidak-ada-kata-ini-xyz", "ganti": "a"},
                                     {"aksi": "ngawur"}])
    assert not dst.exists()
    msg = str(e.value)
    assert "P9999" in msg and "tidak ditemukan" in msg and "tidak dikenal" in msg and "Tidak ada yang diubah" in msg


def test_outline_and_revise_pptx(pptx_file, tmp_path):
    o = outline_pptx(pptx_file)
    assert "--- Slide 1 ---" in o and "S1.sh" in o
    prs = Presentation(str(pptx_file))
    # cari bentuk teks pada slide 1 dan sel tabel/grafik bila ada
    tf_id = next(l.split()[0] for l in o.splitlines() if l.startswith("S1.sh") and "[Tabel" not in l and "[Grafik" not in l)
    dst = tmp_path / "d_rev.pptx"
    rep = revise_pptx(pptx_file, dst, [{"aksi": "ganti_shape_teks", "id": tf_id, "teks": "Judul Revisi"}])
    assert "Judul Revisi" in outline_pptx(dst) and len(rep) == 1
    assert "Judul Revisi" not in outline_pptx(pptx_file)


def test_pptx_table_cell_and_chart_data(pptx_file, tmp_path):
    o = outline_pptx(pptx_file)
    tbl = next((l.split()[0] for l in o.splitlines() if "[Tabel" in l), None)
    chart = next((l.split()[0] for l in o.splitlines() if "[Grafik" in l), None)
    ops = []
    if tbl:
        ops.append({"aksi": "ganti_sel", "id": f"{tbl}.r2c2", "teks": "999"})
    if chart:
        prs = Presentation(str(pptx_file))
        cats = None
        for s in prs.slides:
            for sh in s.shapes:
                if getattr(sh, "has_chart", False) and sh.has_chart:
                    cats = list(sh.chart.plots[0].categories)
                    break
        ops.append({"aksi": "ganti_data_grafik", "id": chart, "seri": [{"nama": "Baru", "nilai": list(range(len(cats)))}]})
    assert ops, "contoh deck harus punya tabel atau grafik"
    dst = tmp_path / "d2.pptx"
    revise_pptx(pptx_file, dst, ops)
    o2 = outline_pptx(dst)
    if tbl:
        assert "999" in o2
    if chart:
        assert 'seri "Baru"' in o2


def test_pptx_invalid_ops(pptx_file, tmp_path):
    with pytest.raises(ReviseError, match="Tidak ada yang diubah"):
        revise_pptx(pptx_file, tmp_path / "z.pptx", [{"aksi": "ganti_shape_teks", "id": "S99.sh1", "teks": "x"},
                                                      {"aksi": "ganti_data_grafik", "id": "S1.sh1", "seri": []}])
    assert not (tmp_path / "z.pptx").exists()


def test_set_paragraph_text_without_runs():
    d = Document()
    p = d.add_paragraph()
    set_paragraph_text(p, "baru")
    assert p.text == "baru"
